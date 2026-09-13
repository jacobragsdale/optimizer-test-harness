"""Excel in and out: dump a BA workbook as text for the compile skill; render manifests as sheets for BAs."""

from __future__ import annotations

from collections.abc import Sequence
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from harness import HarnessError
from harness.run import Manifest

RED = PatternFill("solid", fgColor="FFC7CE")
RESULT_COLUMNS = [
    "case_id",
    "title",
    "started_at_utc",
    "verdict",
    "interpretation",
    "portfolios",
    "run_id",
    "run_url",
    "final_status",
    "expected_outcome",
    "evaluation_plan",
    "evidence",
    "notes",
    "changes",
    "all_reverted",
    "errors",
    "manifest",
]
RUN_COLUMNS = ["case_id", "run_id", "run_url", "submitted_at_utc", "final_status", "succeeded", "output_path", "status_history"]
CHANGE_COLUMNS = ["case_id", "table", "key", "column", "before", "after", "reverted", "revert_ok", "drifted", "observed_before_revert", "note"]


def dump_workbook(path: Path) -> str:
    """Every non-empty row of every sheet, as pipe-separated text with the original row number."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    out: list[str] = []
    for sheet in workbook.worksheets:
        out.append(f"## sheet: {sheet.title}")
        for row in sheet.iter_rows():
            values = ["" if cell.value is None else str(cell.value).replace("\n", " / ") for cell in row]
            if any(values):
                out.append(f"| r{row[0].row} | " + " | ".join(values) + " |")
        out.append("")
    return "\n".join(out)


def render(manifests: Sequence[tuple[Path, Manifest]], out: Path, workbook: Path | None) -> None:
    """Write Results, Runs and Changes sheets. With `workbook`, they are appended to a copy of the BA's file."""
    if workbook is not None and workbook.resolve() == out.resolve():
        msg = "refusing to overwrite the BA workbook; pass a different --out"
        raise HarnessError(msg)
    book = Workbook() if workbook is None else load_workbook(workbook)
    if workbook is None and book.active is not None:
        book.remove(book.active)
    results, runs, changes = (_sheet(book, title, cols) for title, cols in (("Results", RESULT_COLUMNS), ("Runs", RUN_COLUMNS), ("Changes", CHANGE_COLUMNS)))
    for path, m in manifests:
        run = m.run
        verdict = m.verdict
        results.append(
            [
                m.case_id,
                m.spec.title,
                _ts(m.started_at),
                verdict.result if verdict else "not recorded",
                m.spec.interpretation.strip(),
                ", ".join(m.spec.portfolio.ids),
                run.run_id if run else "",
                run.url if run else "",
                run.final_status if run else "",
                m.spec.expected.outcome,
                "\n".join(m.spec.expected.evaluation_plan),
                "\n".join(verdict.evidence) if verdict else "",
                verdict.notes if verdict else "",
                len(m.ledger),
                m.all_reverted,
                "\n".join(m.errors),
                str(path),
            ]
        )
        if run:
            _link(results, RESULT_COLUMNS.index("run_url") + 1, run.url)
            runs.append([m.case_id, run.run_id, run.url, _ts(run.submitted_at), run.final_status, run.succeeded, run.output_path, "\n".join(f"{_ts(e.at)} {e.status}" for e in run.status_history)])
            _link(runs, RUN_COLUMNS.index("run_url") + 1, run.url)
        for e in m.ledger:
            changes.append([m.case_id, e.table, str(e.key), e.column, str(e.before), str(e.after), e.reverted, e.revert_ok, e.drifted, str(e.observed_before_revert), e.note])
            if not e.revert_ok:
                for cell in changes[changes.max_row]:
                    cell.fill = RED
    book.save(out)


def _sheet(book: Workbook, title: str, columns: Sequence[str]) -> Worksheet:
    if title in book.sheetnames:
        book.remove(book[title])
    sheet: Worksheet = book.create_sheet(title)
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    return sheet


def _link(sheet: Worksheet, column: int, url: str) -> None:
    cell = sheet.cell(row=sheet.max_row, column=column, value=url)
    cell.hyperlink = url
    cell.style = "Hyperlink"


def _ts(value: datetime) -> str:
    return value.isoformat(timespec="seconds")
