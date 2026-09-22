"""End-to-end checks against sqlite and the demo app: set up, run, collect, revert, refuse, report.

Only template-owned files are used (demo/, harness/), so these pass unchanged in every app stamped from the template.
"""

from __future__ import annotations

import sqlite3
import threading
from collections.abc import Callable, Iterator
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from typing import cast

import pytest
from openpyxl import load_workbook

from harness import HarnessError
from harness.api import Http, urllib_session
from harness.cli import main
from harness.config import Config, PathsConfig, Settings, load_config
from harness.db import Connection, LedgerEntry, bind, jsonable, query
from harness.report import render
from harness.results import data_files, open_views
from harness.results import query as data_query
from harness.run import MARKER, Manifest, load_manifest, output_paths, pending_revert, record_verdict, revert_ledger, run_case
from harness.spec import Spec, check_spec, load_spec
from harness.stub import DemoApp, init_sample_db

DEMO = Path(__file__).parents[1] / "demo"
NO_SLEEP = lambda _seconds: None  # noqa: E731 -- a named constant reads better than a def here


@pytest.fixture
def db(tmp_path: Path) -> Path:
    path = tmp_path / "qa.sqlite"
    init_sample_db(path)
    return path


@pytest.fixture
def app(tmp_path: Path, db: Path) -> Iterator[DemoApp]:
    server = DemoApp(("127.0.0.1", 0), db, tmp_path / "app-output")
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield server
    server.shutdown()
    server.server_close()


@pytest.fixture
def config(app: DemoApp, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Config:
    monkeypatch.setenv("HARNESS_APP_PASSWORD", "demo")
    demo = load_config(DEMO / "harness.toml")
    return demo.model_copy(update={"api": demo.api.model_copy(update={"base_url": f"http://127.0.0.1:{app.port}"}), "paths": PathsConfig(runs=str(tmp_path / "runs"))})


def demo_spec(case: str, **overrides: object) -> Spec:
    spec = load_spec(DEMO / "specs" / f"{case}.yaml")
    return Spec.model_validate(spec.model_dump() | overrides) if overrides else spec


def run(spec: Spec, config: Config, db: Path, *, http: Http | None = None, connect_db: Callable[[], Connection] | None = None) -> tuple[Manifest, Path]:
    return run_case(spec, config, Settings(db_dsn=str(db)), Path(config.paths.runs), sleep=NO_SLEEP, http=http, connect_db=connect_db)


def snapshot(db: Path) -> list[tuple[object, ...]]:
    """Everything a test may change (app-written JOB_RESULT rows are a known leftover, so excluded)."""
    conn = sqlite3.connect(db)
    try:
        return [tuple(row) for table in ("DISCOUNT_RULE", "PRODUCT") for row in conn.execute(f"SELECT * FROM {table} ORDER BY 1")]
    finally:
        conn.close()


def evidence(config: Config, manifest_file: Path, sql: str) -> list[tuple[object, ...]]:
    manifest = load_manifest(manifest_file)
    con, _, _ = open_views(data_files([manifest_file.parent, *output_paths(manifest, config)]))
    try:
        return data_query(con, sql)[1]
    finally:
        con.close()


PRICES = "SELECT p.product_id, p.final_price FROM (SELECT unnest(prices) AS p FROM get_prices)"


@pytest.mark.parametrize(
    ("case", "prices"), [("EXAMPLE-001", [("P-200", 175.0)]), ("EXAMPLE-002", [("P-400", 72.0)]), ("EXAMPLE-003", [("P-100", 40.0)]), ("EXAMPLE-004", [("P-300", 660.0)]), ("EXAMPLE-005", None)]
)
def test_demo_cases_run_to_their_expected_outcome_and_revert(config: Config, db: Path, case: str, prices: list[tuple[str, float]] | None) -> None:
    fresh = snapshot(db)
    manifest, path = run(demo_spec(case), config, db)

    assert manifest.errors == []
    assert manifest.run_ok is (prices is not None)
    assert manifest.outcome == ("done" if prices is not None else "failed")
    assert manifest.all_reverted
    assert snapshot(db) == fresh
    assert pending_revert(Path(config.paths.runs)) is None
    assert manifest.link == f"{config.api.base_url}/ui/jobs/{manifest.vars['job_id']}"
    if prices is not None:
        assert evidence(config, path, PRICES) == prices
    else:
        assert (manifest.outcome, [s.op for s in manifest.steps if s.phase == "run"]) == ("failed", ["submit_job", "wait_job"])


def test_collect_snapshots_the_database_before_the_revert(config: Config, db: Path) -> None:
    _, path = run(demo_spec("EXAMPLE-001"), config, db)

    assert evidence(config, path, "SELECT RULE_ID, PERCENT FROM rule_during_run") == [("R-2", 30.0)]
    assert evidence(config, path, "SELECT PRODUCT_ID, FINAL_PRICE FROM job_results") == [("P-200", 175.0)]
    assert snapshot(db)[1][2] == 15.0  # R-2 is back to its original percent


def test_app_output_files_join_the_evidence(config: Config, db: Path) -> None:
    _, path = run(demo_spec("EXAMPLE-004"), config, db)

    assert evidence(config, path, "SELECT product_id, markup_percent, final_price FROM prices_export") == [("P-300", 10.0, 660.0)]


def test_api_setup_is_undone_and_the_password_never_reaches_disk(config: Config, db: Path, app: DemoApp) -> None:
    manifest, path = run(demo_spec("EXAMPLE-004"), config, db)

    assert app.price_lists == {}
    assert [(s.phase, s.op, s.ok) for s in manifest.steps if s.phase in {"setup", "undo"}] == [("setup", "create_price_list", True), ("undo", "delete_price_list", True)]
    assert "${HARNESS_APP_PASSWORD}" in path.read_text(encoding="utf-8")
    assert "demo" not in [s.request_body.get("password") for s in manifest.steps if isinstance(s.request_body, dict)]


def test_a_run_step_rejected_by_the_app_is_a_result_not_an_error(config: Config, db: Path) -> None:
    unknown_list = demo_spec("EXAMPLE-001", run=[{"api": "submit_job", "body": {"product_ids": ["P-200"], "price_list_id": "PL-nope"}}], collect=[])
    manifest, _ = run(unknown_list, config, db)

    assert manifest.errors == []
    assert (manifest.run_ok, manifest.outcome) == (False, "HTTP 404")
    assert manifest.all_reverted


def test_everything_reverts_when_the_app_breaks_mid_case(config: Config, db: Path, app: DemoApp) -> None:
    real = urllib_session(10)

    def flaky(method: str, url: str, body: bytes | None, headers: dict[str, str]) -> tuple[int, str, bytes]:
        if url.endswith("/jobs") and method == "POST":
            msg = "connection reset"
            raise HarnessError(msg)
        return real(method, url, body, headers)

    both = demo_spec("EXAMPLE-004", setup=[{"db": "delete", "table": "DISCOUNT_RULE", "key": {"RULE_ID": "R-1"}}, {"api": "create_price_list", "body": {"markup_percent": 5}}])
    fresh = snapshot(db)
    manifest, _ = run(both, config, db, http=cast("Http", flaky))

    assert manifest.errors == ["HarnessError: connection reset"]
    assert manifest.all_reverted
    assert snapshot(db) == fresh
    assert app.price_lists == {}


def test_before_images_reach_disk_before_the_changes_commit(config: Config, db: Path) -> None:
    runs = Path(config.paths.runs)
    ledgers_at_commit: list[int] = []

    class Spy:
        def __init__(self) -> None:
            self.conn = sqlite3.connect(db)

        def cursor(self) -> sqlite3.Cursor:
            return self.conn.cursor()

        def commit(self) -> None:
            ledgers_at_commit.append(len(load_manifest(next(runs.rglob("manifest.json"))).ledger))
            self.conn.commit()

        def rollback(self) -> None:
            self.conn.rollback()

        def close(self) -> None:
            self.conn.close()

    manifest, _ = run(demo_spec("EXAMPLE-001"), config, db, connect_db=lambda: cast("Connection", Spy()))

    assert manifest.errors == []
    assert ledgers_at_commit[0] == 1  # the apply commit


def test_an_interrupted_api_call_is_reported_not_guessed(config: Config) -> None:
    interrupted = LedgerEntry(kind="api", op="create_price_list", undo="delete_price_list")

    entries, errors = revert_ledger([interrupted], config, None, lambda _op, _vars: pytest.fail("must not call the undo"))

    assert entries[0].revert_ok is False
    assert "clean up by hand" in errors[0]


@pytest.mark.parametrize(("value", "driver_type"), [(datetime(2026, 9, 1, 8, 30, tzinfo=UTC), datetime), (date(2026, 9, 1), date), (Decimal("0.10"), Decimal), (b"\x00\x01", bytes)])
def test_before_images_keep_their_driver_type(value: object, driver_type: type) -> None:
    restored = bind(jsonable(value))

    assert (restored, type(restored)) == (value, driver_type)


@pytest.mark.parametrize(
    ("setup", "problem"),
    [
        ({"db": "update", "table": "PRODUCT", "key": {"PRODUCT_ID": "P-200"}, "set": {"BASE_PRICE": 1}}, "not in the [[allow]] list"),
        ({"db": "update", "table": "DISCOUNT_RULE", "key": {"RULE_ID": "R-2"}, "set": {"RULE_ID": "x"}}, "not allowed"),
        ({"db": "update", "table": "DISCOUNT_RULE", "key": {"PRODUCT_ID": "P-200"}, "set": {"PERCENT": 1}}, "key must be exactly RULE_ID"),
        ({"db": "insert", "table": "DISCOUNT_RULE", "values": {"PRODUCT_ID": "P-200", "PERCENT": 1}}, "must include the key"),
        ({"api": "submit_job", "body": {"product_ids": ["P-200"]}}, "has no undo"),
        ({"api": "wait_job"}, "no earlier step captures"),
    ],
    ids=["table", "column", "partial key", "insert without key", "setup op without undo", "variable not captured yet"],
)
def test_spec_refusals(config: Config, setup: dict[str, object], problem: str) -> None:
    problems = check_spec(demo_spec("EXAMPLE-001", setup=[setup]), config)

    assert any(problem in p for p in problems), problems


def test_spec_must_mention_its_subject_and_collect_only_reads(config: Config) -> None:
    spec = demo_spec("EXAMPLE-001", subject={"criteria": "c", "ids": ["P-999"], "rationale": "r"}, collect=[{"name": "x", "sql": "DELETE FROM JOB_RESULT"}])

    assert check_spec(spec, config) == [
        "collect[1] x: only a single read-only SELECT/WITH statement runs here (found 'DELETE'; if it is inside a string literal, rephrase); data changes go through a spec",
        "subject 'P-999' is in subject.ids but no setup, run or collect step mentions it",
    ]


@pytest.mark.parametrize(
    ("setup", "message"),
    [
        ([{"db": "update", "table": "DISCOUNT_RULE", "key": {"RULE_ID": "R-9"}, "set": {"PERCENT": 1}}], "expected exactly 1 row, found 0"),
        ([{"db": "insert", "table": "DISCOUNT_RULE", "values": {"RULE_ID": "R-1", "PERCENT": 1}}], "row already exists"),
    ],
    ids=["no such row", "insert over an existing row"],
)
def test_a_setup_the_database_refuses_changes_nothing(config: Config, db: Path, setup: list[dict[str, object]], message: str) -> None:
    fresh = snapshot(db)
    manifest, _ = run(demo_spec("EXAMPLE-001", setup=setup), config, db)

    assert manifest.ledger == []
    assert message in manifest.errors[0]
    assert snapshot(db) == fresh
    assert pending_revert(Path(config.paths.runs)) is None


def test_refuses_an_api_that_does_not_look_like_qa(config: Config, db: Path) -> None:
    guarded = config.model_copy(update={"db": config.db.model_copy(update={"qa_markers": ["qa"]})})

    with pytest.raises(HarnessError, match=r"api\.base_url"):
        run_case(demo_spec("EXAMPLE-001"), guarded, Settings(db_dsn=str(db) + "-qa"), Path(config.paths.runs))


def test_revert_clears_a_marker_left_by_a_run_killed_before_it_changed_anything(config: Config, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    runs = Path(config.paths.runs)
    runs.mkdir(parents=True)
    ghost = (runs / "EXAMPLE-001" / "never-written" / "manifest.json").resolve()
    (runs / MARKER).write_text(str(ghost), encoding="utf-8")
    (tmp_path / "harness.toml").write_text(_toml(config), encoding="utf-8")

    code = main(["--config", str(tmp_path / "harness.toml"), "revert", str(ghost)])

    assert code == 0
    assert "cleared the PENDING_REVERT marker" in capsys.readouterr().out
    assert pending_revert(runs) is None


@pytest.mark.parametrize(
    "sql",
    ["WITH x AS (SELECT * FROM PRODUCT) DELETE FROM PRODUCT", "SELECT * INTO copy FROM PRODUCT", "SELECT 1; DROP TABLE PRODUCT", "UPDATE PRODUCT SET NAME = 'x'"],
    ids=["cte delete", "select into", "two statements", "update"],
)
def test_sql_refuses_anything_that_writes(db: Path, sql: str) -> None:
    conn = cast("Connection", sqlite3.connect(db))
    try:
        with pytest.raises(HarnessError, match="read-only"):
            query(conn, sql, 10)
        assert query(conn, "SELECT PRODUCT_ID FROM PRODUCT WHERE NAME LIKE '%Desk%' ORDER BY 1", 10)[1] == [("P-100",), ("P-300",)]
    finally:
        conn.close()


def test_report_renders_results_runs_and_changes(config: Config, db: Path, tmp_path: Path) -> None:
    _, path = run(demo_spec("EXAMPLE-001"), config, db)
    manifest = record_verdict(path, "pass", ["P-200 final_price 175.0"], "discount applied")
    out = tmp_path / "results.xlsx"

    render([(path, manifest)], out, None)

    book = load_workbook(out)
    assert book.sheetnames == ["Results", "Runs", "Changes"]
    results = list(book["Results"].iter_rows(values_only=True))
    assert results[0][:4] == ("case_id", "title", "started_at_utc", "verdict")
    assert results[1][:4] == ("EXAMPLE-001", "A bigger chair discount lowers the chair price", manifest.started_at.isoformat(timespec="seconds"), "pass")
    assert book["Results"]["G2"].hyperlink.target == manifest.link
    assert [r[2] for r in book["Runs"].iter_rows(min_row=2, values_only=True)] == ["login", "submit_job", "wait_job", "get_prices"]
    assert next(iter(book["Changes"].iter_rows(min_row=2, values_only=True)))[1:7] == ("db", "update", "DISCOUNT_RULE", '{"RULE_ID": "R-2"}', '{"PERCENT": 15.0}', '{"PERCENT": 30}')


def test_cli_plan_changes_nothing_and_exits_zero(config: Config, db: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]) -> None:
    (tmp_path / "harness.toml").write_text(_toml(config), encoding="utf-8")
    monkeypatch.setenv("HARNESS_DB_DSN", str(db))
    fresh = snapshot(db)

    code = main(["--config", str(tmp_path / "harness.toml"), "run-case", str(DEMO / "specs" / "EXAMPLE-003.yaml")])

    assert code == 0
    assert "db delete DISCOUNT_RULE {'RULE_ID': 'R-1'}: deletes" in capsys.readouterr().out
    assert snapshot(db) == fresh


def _toml(config: Config) -> str:
    """The demo config with this test's base_url and runs path."""
    text = (DEMO / "harness.toml").read_text(encoding="utf-8")
    text = text.replace('base_url = "http://127.0.0.1:8765"', f'base_url = "{config.api.base_url}"').replace('runs = "runs"', f'runs = "{config.paths.runs}"')
    assert f'"{config.paths.runs}"' in text
    return text
