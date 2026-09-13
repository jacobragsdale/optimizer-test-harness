"""End-to-end checks against sqlite and the stub optimizer: apply, run, revert, refuse, report."""

from __future__ import annotations

import sqlite3
import threading
from collections.abc import Iterator
from pathlib import Path

import pytest
from openpyxl import load_workbook

from harness import HarnessError
from harness.cli import main
from harness.config import Config, Settings
from harness.report import render
from harness.run import load_manifest, pending_revert, record_verdict, run_case
from harness.spec import Spec
from harness.stub import StubOptimizer, init_sample_db

NO_SLEEP = lambda _seconds: None  # noqa: E731 -- a named constant reads better than a def here


@pytest.fixture
def stub(tmp_path: Path) -> Iterator[StubOptimizer]:
    server = StubOptimizer(("127.0.0.1", 0), tmp_path / "optimizer-output")
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield server
    server.shutdown()
    server.server_close()


@pytest.fixture
def config(stub: StubOptimizer, tmp_path: Path) -> Config:
    base = f"http://127.0.0.1:{stub.port}"
    return Config.model_validate(
        {
            "db": {"dialect": "sqlite", "qa_markers": []},
            "allow": [{"table": "PORTFOLIO_CONSTRAINT", "key": ["PORTFOLIO_ID", "CONSTRAINT_TYPE"], "columns": ["LIMIT_VALUE", "ENABLED"]}],
            "api": {
                "submit_url": f"{base}/runs",
                "status_url": f"{base}/runs/{{run_id}}",
                "run_url": f"{base}/ui/runs/{{run_id}}",
                "run_id_field": "run_id",
                "status_field": "status",
                "output_path_field": "output_path",
                "terminal_statuses": ["done", "failed"],
                "success_statuses": ["done"],
                "poll_seconds": 0.0,
                "timeout_seconds": 30.0,
            },
            "paths": {"runs": str(tmp_path / "runs")},
        }
    )


@pytest.fixture
def db(tmp_path: Path) -> Path:
    path = tmp_path / "qa.sqlite"
    init_sample_db(path)
    return path


def spec(**overrides: object) -> Spec:
    base: dict[str, object] = {
        "case_id": "TC-001",
        "title": "Sector cap binds",
        "source": {"sheet": "New Feature", "row": 4},
        "description": "Tighten the sector cap on a mid cap portfolio and confirm no sector breaches it.",
        "interpretation": "Lower SECTOR_MAX on PF-1002 from 0.20 to 0.10 and check the holdings output.",
        "portfolio": {"criteria": "mid cap with a sector constraint", "ids": ["PF-1002"], "rationale": "only mid cap portfolio with SECTOR_MAX enabled"},
        "changes": [{"table": "PORTFOLIO_CONSTRAINT", "key": {"PORTFOLIO_ID": "PF-1002", "CONSTRAINT_TYPE": "SECTOR_MAX"}, "column": "LIMIT_VALUE", "value": 0.1}],
        "run": {"params": {"portfolio_id": "PF-1002"}},
        "expected": {"outcome": "success", "text": "No sector above 10%", "evaluation_plan": ["sum weight by sector in holdings; max must be <= 0.10"]},
    }
    return Spec.model_validate(base | overrides)


def limit_value(db: Path, portfolio: str = "PF-1002", constraint: str = "SECTOR_MAX") -> object:
    conn = sqlite3.connect(db)
    try:
        return conn.execute("SELECT LIMIT_VALUE FROM PORTFOLIO_CONSTRAINT WHERE PORTFOLIO_ID=? AND CONSTRAINT_TYPE=?", (portfolio, constraint)).fetchone()[0]
    finally:
        conn.close()


def test_run_case_applies_runs_and_reverts(config: Config, db: Path) -> None:
    runs = Path(config.paths.runs)
    manifest, path = run_case(spec(), config, Settings(db_dsn=str(db)), runs, sleep=NO_SLEEP)

    assert manifest.errors == []
    assert manifest.run is not None
    assert manifest.run.final_status == "done"
    assert manifest.run.succeeded is True
    assert [e.status for e in manifest.run.status_history] == ["running", "done"]
    assert manifest.run.output_path is not None
    assert (Path(manifest.run.output_path) / "holdings.parquet").exists()
    assert [(e.before, e.after, e.reverted, e.revert_ok, e.drifted) for e in manifest.ledger] == [(0.2, 0.1, True, True, False)]
    assert limit_value(db) == 0.2
    assert pending_revert(runs) is None
    reloaded = load_manifest(path)
    assert reloaded.run is not None
    assert reloaded.run.run_id == manifest.run.run_id


def test_revert_still_runs_when_submit_fails(config: Config, db: Path) -> None:
    broken = config.model_copy(update={"api": config.api.model_copy(update={"submit_url": config.api.submit_url + "/nope"})})
    manifest, _ = run_case(spec(), broken, Settings(db_dsn=str(db)), Path(config.paths.runs), sleep=NO_SLEEP)

    assert len(manifest.errors) == 1
    assert "HTTP 404" in manifest.errors[0]
    assert manifest.run is None
    assert manifest.all_reverted
    assert limit_value(db) == 0.2
    assert pending_revert(Path(config.paths.runs)) is None


def test_failed_optimizer_status_is_a_result_not_an_error(config: Config, db: Path) -> None:
    negative = spec(run={"params": {"portfolio_id": "PF-1002", "fail": True}}, expected={"outcome": "failure", "text": "rejected as infeasible", "evaluation_plan": ["status is failed"]})
    manifest, _ = run_case(negative, config, Settings(db_dsn=str(db)), Path(config.paths.runs), sleep=NO_SLEEP)

    assert manifest.errors == []
    assert manifest.run is not None
    assert (manifest.run.final_status, manifest.run.succeeded, manifest.run.output_path) == ("failed", False, None)
    assert limit_value(db) == 0.2


@pytest.mark.parametrize(
    "change",
    [
        {"table": "PORTFOLIO", "key": {"PORTFOLIO_ID": "PF-1002"}, "column": "NAME", "value": "x"},
        {"table": "PORTFOLIO_CONSTRAINT", "key": {"PORTFOLIO_ID": "PF-1002", "CONSTRAINT_TYPE": "SECTOR_MAX"}, "column": "PORTFOLIO_ID", "value": "x"},
        {"table": "PORTFOLIO_CONSTRAINT", "key": {"PORTFOLIO_ID": "PF-1002"}, "column": "LIMIT_VALUE", "value": 0.1},
        {"table": "PORTFOLIO_CONSTRAINT", "key": {"PORTFOLIO_ID": "PF-1002", "CONSTRAINT_TYPE": "MISSING"}, "column": "LIMIT_VALUE", "value": 0.1},
    ],
    ids=["table not allowed", "column not allowed", "partial key", "no such row"],
)
def test_refuses_changes_outside_the_allowlist_or_off_one_row(config: Config, db: Path, change: dict[str, object]) -> None:
    runs = Path(config.paths.runs)
    try:
        manifest, _ = run_case(spec(changes=[change]), config, Settings(db_dsn=str(db)), runs, sleep=NO_SLEEP)
    except HarnessError:
        manifest = None
    if manifest is not None:  # the no-such-row case is refused at apply time, inside the transaction
        assert manifest.ledger == []
        assert manifest.errors and "expected exactly 1 row" in manifest.errors[0]
    assert limit_value(db) == 0.2
    assert pending_revert(runs) is None


def test_report_renders_results_runs_and_changes(config: Config, db: Path, tmp_path: Path) -> None:
    _, path = run_case(spec(), config, Settings(db_dsn=str(db)), Path(config.paths.runs), sleep=NO_SLEEP)
    manifest = record_verdict(path, "pass", ["max sector weight 0.0987"], "within cap")
    out = tmp_path / "results.xlsx"

    render([(path, manifest)], out, None)

    book = load_workbook(out)
    assert book.sheetnames == ["Results", "Runs", "Changes"]
    results = list(book["Results"].iter_rows(values_only=True))
    assert results[0][:4] == ("case_id", "title", "started_at_utc", "verdict")
    assert results[1][:4] == ("TC-001", "Sector cap binds", manifest.started_at.isoformat(timespec="seconds"), "pass")
    assert manifest.run is not None
    assert book["Runs"]["C2"].hyperlink.target == manifest.run.url
    assert next(iter(book["Changes"].iter_rows(min_row=2, values_only=True)))[1:6] == (
        "PORTFOLIO_CONSTRAINT",
        "{'PORTFOLIO_ID': 'PF-1002', 'CONSTRAINT_TYPE': 'SECTOR_MAX'}",
        "LIMIT_VALUE",
        "0.2",
        "0.1",
    )


def test_cli_plan_changes_nothing_and_exits_zero(config: Config, db: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]) -> None:
    import yaml

    (tmp_path / "harness.toml").write_text(_toml(config), encoding="utf-8")
    (tmp_path / "spec.yaml").write_text(yaml.safe_dump(spec().model_dump()), encoding="utf-8")
    monkeypatch.setenv("HARNESS_DB_DSN", str(db))

    code = main(["--config", str(tmp_path / "harness.toml"), "run-case", str(tmp_path / "spec.yaml")])

    assert code == 0
    assert "0.2 -> 0.1" in capsys.readouterr().out
    assert limit_value(db) == 0.2


def _toml(config: Config) -> str:
    import tomllib

    text = f"""
[db]
dialect = "sqlite"
qa_markers = []

[[allow]]
table = "PORTFOLIO_CONSTRAINT"
key = ["PORTFOLIO_ID", "CONSTRAINT_TYPE"]
columns = ["LIMIT_VALUE", "ENABLED"]

[api]
submit_url = "{config.api.submit_url}"
status_url = "{config.api.status_url}"
run_url = "{config.api.run_url}"
run_id_field = "run_id"
status_field = "status"
output_path_field = "output_path"
terminal_statuses = ["done", "failed"]
success_statuses = ["done"]
poll_seconds = 0.0
timeout_seconds = 30.0

[paths]
runs = "{config.paths.runs}"
"""
    assert Config.model_validate(tomllib.loads(text)) == config
    return text
