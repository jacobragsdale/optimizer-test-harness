"""Excel in and out: dump a BA workbook as text for the compile skill; render manifests as sheets for BAs."""

from __future__ import annotations

import json
from collections.abc import Sequence
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from harness import HarnessError
from harness.db import LedgerEntry, plain
from harness.run import Manifest

RED = PatternFill("solid", fgColor="FFC7CE")
RESULT_COLUMNS = [
    "case_id",
    "title",
    "started_at_utc",
    "verdict",
    "interpretation",
    "subjects",
    "link",
    "outcome",
    "expected_outcome",
    "evaluation_plan",
    "evidence",
    "notes",
    "setup_steps",
    "all_reverted",
    "errors",
    "manifest",
]
RUN_COLUMNS = ["case_id", "phase", "op", "method", "url", "http_status", "final_status", "ok", "started_at_utc", "status_history", "response_file", "detail"]
CHANGE_COLUMNS = ["case_id", "kind", "op", "target", "key_or_vars", "before", "after", "reverted", "revert_ok", "drifted", "observed_before_revert", "note"]


def dump_workbook(path: Path) -> str:
    """Every non-empty row of every sheet, as pipe-separated text with the original row number."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    out: list[str] = []
    for sheet in workbook.worksheets:
        out.append(f"## sheet: {sheet.title}")
        for row in sheet.iter_rows():
            values = ["" if cell.value is None else str(cell.value).replace("\n", " / ") for cell in row]
            if any(values):
                out.append(f"| r{row[0].row} | " + " | ".join(values) + " |")
        out.append("")
    return "\n".join(out)


def render(manifests: Sequence[tuple[Path, Manifest]], out: Path, workbook: Path | None) -> None:
    """Write Results, Runs and Changes sheets. With `workbook`, they are appended to a copy of the BA's file."""
    if workbook is not None and workbook.resolve() == out.resolve():
        msg = "refusing to overwrite the BA workbook; pass a different --out"
        raise HarnessError(msg)
    book = Workbook() if workbook is None else load_workbook(workbook)
    if workbook is None and book.active is not None:
        book.remove(book.active)
    results, runs, changes = (_sheet(book, title, cols) for title, cols in (("Results", RESULT_COLUMNS), ("Runs", RUN_COLUMNS), ("Changes", CHANGE_COLUMNS)))
    for path, m in manifests:
        verdict = m.verdict
        results.append(
            [
                m.case_id,
                m.spec.title,
                _ts(m.started_at),
                verdict.result if verdict else "not recorded",
                m.spec.interpretation.strip(),
                ", ".join(m.spec.subject.ids),
                m.link or "",
                m.outcome,
                m.spec.expected.outcome,
                "\n".join(m.spec.expected.evaluation_plan),
                "\n".join(verdict.evidence) if verdict else "",
                verdict.notes if verdict else "",
                len(m.ledger),
                m.all_reverted,
                "\n".join(m.errors),
                str(path),
            ]
        )
        if m.link:
            _link(results, RESULT_COLUMNS.index("link") + 1, m.link)
        for s in m.steps:
            history = "\n".join(f"{_ts(e.at)} {e.status}" for e in s.status_history)
            runs.append([m.case_id, s.phase, s.op, s.method, s.url, s.http_status, s.final_status, s.ok, _ts(s.started_at), history, s.response_file, s.detail])
        for e in m.ledger:
            changes.append([m.case_id, *_change_row(e)])
            if not e.done:
                for cell in changes[changes.max_row]:
                    cell.fill = RED
    book.save(out)


def _change_row(e: LedgerEntry) -> list[object]:
    if e.kind == "api":
        return ["api", e.op, f"undo: {e.undo}", _text(e.vars), _text(e.request), f"HTTP {e.http_status}", e.reverted, e.revert_ok, e.drifted, "", e.note]
    return ["db", e.op, e.table, _text(e.key), _text(e.before), _text(e.after), e.reverted, e.revert_ok, e.drifted, _text(e.observed_before_revert), e.note]


def _text(value: object) -> str:
    if isinstance(value, dict):
        return json.dumps({k: plain(v) for k, v in value.items()}, default=str) if value else ""
    return "" if value is None else json.dumps(value, default=str)


def _sheet(book: Workbook, title: str, columns: Sequence[str]) -> Worksheet:
    if title in book.sheetnames:
        book.remove(book[title])
    sheet: Worksheet = book.create_sheet(title)
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    return sheet


def _link(sheet: Worksheet, column: int, url: str) -> None:
    cell = sheet.cell(row=sheet.max_row, column=column, value=url)
    cell.hyperlink = url
    cell.style = "Hyperlink"


def _ts(value: datetime) -> str:
    return value.isoformat(timespec="seconds")
